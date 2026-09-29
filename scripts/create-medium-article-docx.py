from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "MEDIUM_ARTICLE.md"
OUTPUT = ROOT / "artifacts" / "Walrus_Promise_Medium_Article.docx"

IMAGE_MAP = {
    "tester-01-success.png": ROOT / "docs" / "evidence" / "real-testers" / "tester-01-success.png",
    "tester-02-success.png": ROOT / "docs" / "evidence" / "real-testers" / "tester-02-success.png",
    "tester-03-success-after-fix.png": ROOT / "docs" / "evidence" / "real-testers" / "tester-03-success-after-fix.png",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F5D9B")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def clean_typography(text: str) -> str:
    replacements = {
        "\u2014": " - ",
        "\u2013": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2026": "...",
        "\u2192": "to",
        "\u00b7": "-",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return re.sub(r"\s+", " ", text).strip()


def add_inline(paragraph, text: str) -> None:
    text = clean_typography(text)
    pattern = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        add_hyperlink(paragraph, match.group(1), match.group(2))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_picture(doc: Document, alt: str, url: str) -> None:
    filename = url.rsplit("/", 1)[-1]
    path = IMAGE_MAP.get(filename)
    if not path or not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(6.7))
    caption = doc.add_paragraph(clean_typography(alt))
    caption.style = doc.styles["Caption"]
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(9)
    props = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F4F3")
    props.append(shading)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(8.5)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.autofit = True
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.text = clean_typography(value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "243447")
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = clean_typography(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            if row_index % 2:
                set_cell_shading(cells[index], "F3F6F8")
    doc.add_paragraph()


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, before, after in (
        ("Title", 28, 0, 12),
        ("Subtitle", 14, 0, 18),
        ("Heading 1", 19, 18, 7),
        ("Heading 2", 14, 14, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph("Walrus Promise Technical Article", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph("Your Chatbot Should Remember the Promise Not the Transcript", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    skipped_first_title = False

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue

        if line.startswith("|") and line.endswith("|"):
            values = [value.strip() for value in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                continue
            table_rows.append(values)
            continue
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

        image = re.fullmatch(r"!\[([^\]]+)\]\((https?://[^)]+)\)", line)
        if image:
            add_picture(doc, image.group(1), image.group(2))
            continue

        if line.startswith("# "):
            if not skipped_first_title:
                skipped_first_title = True
            continue
        if line.startswith("## "):
            doc.add_paragraph(clean_typography(line[3:]), style="Heading 1")
            continue
        if line.startswith("### "):
            doc.add_paragraph(clean_typography(line[4:]), style="Heading 2")
            continue
        if line in ("---", "***", "___"):
            continue
        if not line.strip():
            continue

        ordered = re.match(r"^\d+\.\s+(.+)$", line)
        if ordered:
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, ordered.group(1))
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            continue
        if line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            run = paragraph.add_run(clean_typography(line[2:]))
            run.italic = True
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, line)

    if table_rows:
        add_table(doc, table_rows)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Walrus Promise Technical Article"
    doc.core_properties.subject = "Hackathon article about cross session memory with Walrus Memory"
    doc.core_properties.author = "Walrus Promise"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
